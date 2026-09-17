#!/usr/bin/env python3
"""
egSYS JiraView — Gerador do Deck Executivo (CEO & Diretoria) — v0.4.1
Padrão Canônico de 8 Slides Widescreen (16:9) em Dark Glass / Modern Slate.
Destaque absoluto para as melhorias em relação ao painel legado do Jira:
- Trilha da jornada (esteira de ação de 8 etapas + tempo em cada fase + posse da bola)
- Dashboards e observabilidade visual (gráficos interativos + banner de síntese)
- Construtor de filtros e persistência de filtros salvos pelo gestor
- Exportação corporativa (.xlsx, .md, .csv) com numeração # 1 a N
- Rastreabilidade com tarefas técnicas da engenharia de software (PSC)
"""

import os
import sys
import subprocess
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# ==============================================================================
# 1. PALETA DE CORES CANÔNICA (Tema Dark Glass / Modern Slate)
# ==============================================================================
COLOR_BG = RGBColor(8, 17, 31)             # #08111F
COLOR_CARD_DARK = RGBColor(16, 32, 56)     # #102038
COLOR_CARD_BLUE = RGBColor(19, 38, 65)     # #132641
COLOR_BORDER = RGBColor(42, 62, 92)        # #2A3E5C
COLOR_BORDER_CYAN = RGBColor(53, 208, 255)
COLOR_BORDER_LIME = RGBColor(200, 255, 74)

COLOR_TEXT_LIGHT = RGBColor(245, 248, 252) # #F5F8FC
COLOR_TEXT_MUTED = RGBColor(157, 176, 199) # #9DB0C7
COLOR_CYAN = RGBColor(53, 208, 255)        # #35D0FF (Acentos Técnicos)
COLOR_LIME = RGBColor(200, 255, 74)        # #C8FF4A (Big Numbers / ROI)
COLOR_WARN = RGBColor(255, 184, 77)        # #FFB84D (Alerta)
COLOR_RED = RGBColor(255, 92, 92)          # #FF5C5C (Risco Crítico)

FONT_DISPLAY = "Aptos Display"
FONT_BODY = "Aptos"


class JiraViewDeckBuilder:
    def __init__(self, output_pptx="egSYS-JiraView-CEO-Executivo.pptx"):
        self.output_pptx = output_pptx
        self.prs = Presentation()
        self.prs.slide_width = Inches(13.333)
        self.prs.slide_height = Inches(7.50)
        self.blank_layout = self.prs.slide_layouts[6]

    def add_slide(self):
        slide = self.prs.slides.add_slide(self.blank_layout)
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.50))
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLOR_BG
        bg.line.fill.background()
        return slide

    def add_header(self, slide, tag_text, title_text, subtitle_text=None, slide_num_str=None, tag_color=COLOR_CYAN):
        # Tag Superior
        tx_tag = slide.shapes.add_textbox(Inches(0.75), Inches(0.35), Inches(8.5), Inches(0.28))
        tf_t = tx_tag.text_frame
        tf_t.word_wrap = True
        tf_t.margin_left = tf_t.margin_top = tf_t.margin_right = tf_t.margin_bottom = 0
        p_t = tf_t.paragraphs[0]
        r_t = p_t.add_run()
        r_t.text = tag_text.upper()
        r_t.font.name = FONT_BODY
        r_t.font.size = Pt(10.5)
        r_t.font.bold = True
        r_t.font.color.rgb = tag_color

        # Número do Slide
        if slide_num_str:
            tx_num = slide.shapes.add_textbox(Inches(11.8), Inches(0.35), Inches(0.78), Inches(0.28))
            tf_n = tx_num.text_frame
            tf_n.word_wrap = True
            tf_n.margin_left = tf_n.margin_top = tf_n.margin_right = tf_n.margin_bottom = 0
            p_n = tf_n.paragraphs[0]
            p_n.alignment = PP_ALIGN.RIGHT
            r_n = p_n.add_run()
            r_n.text = slide_num_str
            r_n.font.name = FONT_BODY
            r_n.font.size = Pt(11)
            r_n.font.bold = True
            r_n.font.color.rgb = COLOR_TEXT_MUTED

        # Título do Slide
        tx_title = slide.shapes.add_textbox(Inches(0.75), Inches(0.68), Inches(11.83), Inches(0.75))
        tf_title = tx_title.text_frame
        tf_title.word_wrap = True
        tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
        p_title = tf_title.paragraphs[0]
        r_title = p_title.add_run()
        r_title.text = title_text
        r_title.font.name = FONT_DISPLAY
        r_title.font.size = Pt(25.0)
        r_title.font.bold = True
        r_title.font.color.rgb = COLOR_TEXT_LIGHT

        # Subtítulo do Slide
        if subtitle_text:
            tx_sub = slide.shapes.add_textbox(Inches(0.75), Inches(1.48), Inches(11.83), Inches(0.50))
            tf_sub = tx_sub.text_frame
            tf_sub.word_wrap = True
            tf_sub.margin_left = tf_sub.margin_top = tf_sub.margin_right = tf_sub.margin_bottom = 0
            p_sub = tf_sub.paragraphs[0]
            r_sub = p_sub.add_run()
            r_sub.text = subtitle_text
            r_sub.font.name = FONT_BODY
            r_sub.font.size = Pt(13.5)
            r_sub.font.color.rgb = COLOR_TEXT_LIGHT

    def add_footer(self, slide, left_text="egSYS JIRAVIEW", right_text="PAINEL DO CLIENTE JSM / INTELIGÊNCIA OPERACIONAL"):
        rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.75), Inches(7.02), Inches(11.83), Inches(0.01))
        rule.fill.solid()
        rule.fill.fore_color.rgb = COLOR_BORDER
        rule.line.color.rgb = COLOR_BORDER

        tx_box = slide.shapes.add_textbox(Inches(0.75), Inches(7.10), Inches(3.5), Inches(0.3))
        tf = tx_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = left_text
        r.font.name = FONT_BODY
        r.font.size = Pt(8.5)
        r.font.bold = True
        r.font.color.rgb = COLOR_TEXT_MUTED

        tx_box_r = slide.shapes.add_textbox(Inches(7.5), Inches(7.10), Inches(5.08), Inches(0.3))
        tf_r = tx_box_r.text_frame
        tf_r.word_wrap = True
        tf_r.margin_left = tf_r.margin_top = tf_r.margin_right = tf_r.margin_bottom = 0
        p_r = tf_r.paragraphs[0]
        p_r.alignment = PP_ALIGN.RIGHT
        r_r = p_r.add_run()
        r_r.text = right_text
        r_r.font.name = FONT_BODY
        r_r.font.size = Pt(8.0)
        r_r.font.bold = True
        r_r.font.color.rgb = COLOR_TEXT_MUTED

    def add_card(self, slide, left, top, width, height, fill_color=COLOR_CARD_DARK, border_color=COLOR_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        card.fill.solid()
        card.fill.fore_color.rgb = fill_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1)
        return card

    def add_callout(self, slide, left, top, width, height, text, accent_color=COLOR_WARN):
        card = self.add_card(slide, left, top, width, height, fill_color=COLOR_CARD_BLUE, border_color=COLOR_BORDER)
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.2)
        tf.margin_right = Inches(0.2)
        tf.margin_top = Inches(0.08)
        tf.margin_bottom = Inches(0.08)
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = text
        r.font.name = FONT_BODY
        r.font.size = Pt(10.5)
        r.font.bold = True
        r.font.color.rgb = accent_color
        return card

    # ==============================================================================
    # SLIDE 1: CAPA & EXECUTIVE BRIEF
    # ==============================================================================
    def build_slide_1(self):
        slide = self.add_slide()
        self.add_header(
            slide,
            tag_text="egSYS ANÁLISE EXECUTIVA  /  2026",
            title_text="egSYS JIRAVIEW: Transparência, Observabilidade e Escala do Atendimento JSM",
            subtitle_text="Superando o portal legado do Jira: novo painel com gráficos, esteira de 8 etapas, filtros salvos e leitura executiva em ≤ 3s.",
            tag_color=COLOR_CYAN,
        )

        # Card de Destaque Superior
        c_destaque = self.add_card(slide, 0.75, 2.30, 11.83, 1.35, fill_color=COLOR_CARD_BLUE, border_color=COLOR_BORDER_CYAN)
        tf_d = c_destaque.text_frame
        tf_d.word_wrap = True
        tf_d.margin_left = tf_d.margin_right = Inches(0.35)
        tf_d.margin_top = Inches(0.25)
        p_d1 = tf_d.paragraphs[0]
        r_d1 = p_d1.add_run()
        r_d1.text = "⚡ VISIBILIDADE EXECUTIVA ≤ 3s   •   📊 OBSERVABILIDADE VISUAL TOTAL   •   🛡️ ESTEIRA COM POSSE DA BOLA"
        r_d1.font.name = FONT_DISPLAY
        r_d1.font.size = Pt(17.0)
        r_d1.font.bold = True
        r_d1.font.color.rgb = COLOR_LIME

        p_d2 = tf_d.add_paragraph()
        p_d2.space_before = Pt(8)
        r_d2 = p_d2.add_run()
        r_d2.text = "Substituição completa do portal rígido do Jira: filtros salvos pelo gestor, dashboards gráficos e radiografia profunda de tickets em 1 container único."
        r_d2.font.name = FONT_BODY
        r_d2.font.size = Pt(11.5)
        r_d2.font.color.rgb = COLOR_TEXT_MUTED

        # Card de Decisão & Resultados
        c_decisao = self.add_card(slide, 0.75, 3.90, 11.83, 2.70, fill_color=COLOR_CARD_DARK, border_color=COLOR_BORDER)
        tf_dec = c_decisao.text_frame
        tf_dec.word_wrap = True
        tf_dec.margin_left = tf_dec.margin_right = Inches(0.35)
        tf_dec.margin_top = Inches(0.25)

        p_h = tf_dec.paragraphs[0]
        r_h = p_h.add_run()
        r_h.text = "AS GRANDES MELHORIAS FUNCIONAIS FRENTE AO PAINEL LEGADO DO JIRA"
        r_h.font.name = FONT_BODY
        r_h.font.size = Pt(12)
        r_h.font.bold = True
        r_h.font.color.rgb = COLOR_CYAN

        bullets = [
            ("Trilha Visual da Jornada (Esteira de 8 Etapas): ", "Fim do status opaco. O cliente acompanha cada etapa (Triagem ➔ Dev ➔ QA ➔ Homologação), o tempo decorrido em cada fase e a Posse da Bola."),
            ("Dashboard Gráfico & Banner de Síntese Matemática: ", "4 gráficos interativos (status, prioridade, solicitantes e tendência) mais o Banner Executivo com decomposição exata do backlog ativo (25 Ativas / 25 Concluídas)."),
            ("Construtor Visual de Filtros & Salvamento Personalizado: ", "Criação de filtros JQL com multiseleção e capacidade de salvar listagens customizadas por gestor (/filtros), inexistente no Jira padrão."),
            ("Rastreabilidade Técnica e Exportação (# 1 a N): ", "Links 1-click para o Jira oficial, mapeamento de tarefas da engenharia (PSC) e exportação em Excel nativo (.xlsx), Markdown e CSV numerados."),
        ]

        for b_title, b_desc in bullets:
            p = tf_dec.add_paragraph()
            p.space_before = Pt(7)
            r1 = p.add_run()
            r1.text = "• " + b_title
            r1.font.name = FONT_BODY
            r1.font.size = Pt(11)
            r1.font.bold = True
            r1.font.color.rgb = COLOR_TEXT_LIGHT

            r2 = p.add_run()
            r2.text = b_desc
            r2.font.name = FONT_BODY
            r2.font.size = Pt(10.5)
            r2.font.color.rgb = COLOR_TEXT_MUTED

        self.add_footer(slide)

    # ==============================================================================
    # SLIDE 2: O DESAFIO & GARGALOS DO PAINEL LEGADO DO JIRA
    # ==============================================================================
    def build_slide_2(self):
        slide = self.add_slide()
        self.add_header(
            slide,
            tag_text="O PROBLEMA  /  GARGALOS DO PORTAL LEGADO DO JIRA",
            title_text="O portal nativo do Jira mantinha o cliente cego sobre o fluxo real.",
            subtitle_text="Sem gráficos, sem esteira de ação, sem filtros salvos e com ambiguidade entre tarefas ativas e entregas passadas.",
            slide_num_str="02",
            tag_color=COLOR_WARN,
        )

        riscos = [
            (0.75, 2.35, "🔴 SEM TRILHA DA JORNADA", "Status Estático e Opaco", "O cliente via apenas um rótulo cru ('Em Andamento'). Impossível saber em que etapa o chamado estava (N2, Dev ou QA), há quanto tempo estava parado e de quem era a responsabilidade da ação."),
            (6.85, 2.35, "🔴 SEM DASHBOARDS OU GRÁFICOS", "Ausência Total de Observabilidade", "Nenhum gráfico disponível ao cliente no portal padrão. Apenas tabelas em 'planilha pesada' sem visão de funil, sem divisão por criticidade e sem identificação de gargalos."),
            (0.75, 4.35, "🔴 SEM FILTRAGEM EFETIVA", "Zero Personalização para o Cliente", "Buscas rudimentares que truncavam em 100 itens. O cliente não conseguia filtrar por frentes de produto, não podia cruzar critérios e era impossível salvar listagens personalizadas."),
            (6.85, 4.35, "🔴 AMBIGUIDADE CRÔNICA DE DADOS", "A Falsa Impressão de Atrasos", "O termo genérico 'Abertas' somava itens concluídos há meses com demandas ativas e chamados parados no próprio cliente, gerando atritos desgastantes e cobranças indevidas."),
        ]

        for left, top, tag, title, desc in riscos:
            card = self.add_card(slide, left, top, 5.73, 1.80)
            tf = card.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_right = Inches(0.25)
            tf.margin_top = Inches(0.18)

            p1 = tf.paragraphs[0]
            r1 = p1.add_run()
            r1.text = tag
            r1.font.name = FONT_BODY
            r1.font.size = Pt(10)
            r1.font.bold = True
            r1.font.color.rgb = COLOR_RED

            p2 = tf.add_paragraph()
            p2.space_before = Pt(3)
            r2 = p2.add_run()
            r2.text = title
            r2.font.name = FONT_DISPLAY
            r2.font.size = Pt(13)
            r2.font.bold = True
            r2.font.color.rgb = COLOR_TEXT_LIGHT

            p3 = tf.add_paragraph()
            p3.space_before = Pt(4)
            r3 = p3.add_run()
            r3.text = desc
            r3.font.name = FONT_BODY
            r3.font.size = Pt(9.5)
            r3.font.color.rgb = COLOR_TEXT_MUTED

        self.add_callout(
            slide,
            0.75, 6.35, 11.83, 0.48,
            "⚠️ Impacto do Portal Legado: Clientes insatisfeitos sem visibilidade, 90% do tempo perdido em conferências manuais e cobranças injustas à equipe egSYS.",
            accent_color=COLOR_WARN,
        )
        self.add_footer(slide)

    # ==============================================================================
    # SLIDE 3: A VIRADA ARQUITETURAL & AS NOVAS FUNCIONALIDADES
    # ==============================================================================
    def build_slide_3(self):
        slide = self.add_slide()
        self.add_header(
            slide,
            tag_text="A TRANSFORMAÇÃO  /  ARQUITETURA & RECURSOS DO PAINEL",
            title_text="O JiraView entrega a observabilidade visual que o Jira nunca teve.",
            subtitle_text="Todas as camadas de controle reunidas em uma plataforma única, ágil, personalizável e de alto desempenho.",
            slide_num_str="03",
            tag_color=COLOR_CYAN,
        )

        pilares = [
            (0.75, 2.25, 3.75, 4.00, "OBSERVABILIDADE VISUAL TOTAL", "Dashboards & Gráficos Táticos", [
                ("Banner Executivo de Backlog: ", "Equação visual no topo eliminando a confusão de números."),
                ("Funil de 4 Estágios: ", "Cards semânticos categorizados por componente da esteira."),
                ("4 Gráficos Interativos: ", "Barras por status, donut de prioridade, polar e linha temporal."),
                ("Modo Escuro Full-Window: ", "Design Dark Glass sem reflexos, ideal para salas de situação."),
            ]),
            (4.79, 2.25, 3.75, 4.00, "FILTRAGEM AVANÇADA & SALVA", "Construtor JQL & CRUD de Filtros", [
                ("Filtros Salvos pelo Gestor: ", "Módulo completo (/filtros) para salvar e reusar visões em 1 clique."),
                ("Construtor Visual de JQL: ", "Filtre por campo, operador e valor sem saber código Jira."),
                ("Filtro por Áreas de Atendimento: ", "Abas automáticas mapeadas por componentes do negócio."),
                ("Governança Temporal: ", "Janelas de 60d, 90d, 6m, 12m e ano atual sem truncamento."),
            ]),
            (8.83, 2.25, 3.75, 4.00, "A TRILHA DA JORNADA & SRE", "Gaveta Lateral (Drawer) & Traefik", [
                ("Esteira de 8 Etapas: ", "Da Triagem N1 à Homologação do Cliente com Posse da Bola."),
                ("Cálculo de Tempo por Fase: ", "Veja quantos dias o ticket passou em cada status real."),
                ("1 Container Multi-Estado: ", "Configuração declarativa em YAML e Cgroups (512MB RAM)."),
                ("Exportação (# 1 a N): ", "Download imediato em Excel (.xlsx), Markdown e CSV."),
            ]),
        ]

        for left, top, width, height, tag, title, bullets in pilares:
            card = self.add_card(slide, left, top, width, height)
            tf = card.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_right = Inches(0.25)
            tf.margin_top = Inches(0.22)

            p1 = tf.paragraphs[0]
            r1 = p1.add_run()
            r1.text = tag
            r1.font.name = FONT_BODY
            r1.font.size = Pt(10)
            r1.font.bold = True
            r1.font.color.rgb = COLOR_CYAN

            p2 = tf.add_paragraph()
            p2.space_before = Pt(3)
            r2 = p2.add_run()
            r2.text = title
            r2.font.name = FONT_DISPLAY
            r2.font.size = Pt(13.5)
            r2.font.bold = True
            r2.font.color.rgb = COLOR_TEXT_LIGHT

            for b_label, b_text in bullets:
                p = tf.add_paragraph()
                p.space_before = Pt(7)
                rb1 = p.add_run()
                rb1.text = "• " + b_label
                rb1.font.name = FONT_BODY
                rb1.font.size = Pt(10)
                rb1.font.bold = True
                rb1.font.color.rgb = COLOR_TEXT_LIGHT

                rb2 = p.add_run()
                rb2.text = b_text
                rb2.font.name = FONT_BODY
                rb2.font.size = Pt(9.5)
                rb2.font.color.rgb = COLOR_TEXT_MUTED

        self.add_callout(
            slide,
            0.75, 6.38, 11.83, 0.48,
            "🎯 Resultado: O cliente ganha poder de auditoria, filtros sob medida e visualização completa sem depender de analistas da egSYS.",
            accent_color=COLOR_LIME,
        )
        self.add_footer(slide)

    # ==============================================================================
    # SLIDE 4: IMPACTO MENSURADO & RETORNO OPERACIONAL (ROI)
    # ==============================================================================
    def build_slide_4(self):
        slide = self.add_slide()
        self.add_header(
            slide,
            tag_text="IMPACTO MENSURADO  /  RESULTADOS AUDITADOS",
            title_text="O salto operacional do cliente: métricas de observabilidade.",
            subtitle_text="Ganhos comprovados em tempo de análise, precisão de backlog e autonomia de filtragem do cliente.",
            slide_num_str="04",
            tag_color=COLOR_LIME,
        )

        bignumbers = [
            (0.75, 2.25, "≤ 3s", "TEMPO PARA LEITURA DE STATUS", "De 30 a 60 segundos navegando em listas confusas para menos de 3 segundos com o Banner e Funil."),
            (6.85, 2.25, "100%", "PRECISÃO NA CONTA DO BACKLOG", "Fim das disputas de quantidade: 0 Novas + 24 Em Atendimento + 1 Aguardando Cliente = 25 Ativas (ciclo 90d)."),
            (0.75, 4.30, "8 Fases", "TRILHA COMPLETA DA JORNADA", "Cada ticket tem sua esteira exibida com Posse da Bola e dias acumulados em cada estágio."),
            (6.85, 4.30, "1-Click", "EXPORTAÇÃO & FILTROS SALVOS", "Gestor salva suas visões favoritas e baixa planilhas formatadas (#: 1 a N) com apenas um clique."),
        ]

        for left, top, num, title, desc in bignumbers:
            card = self.add_card(slide, left, top, 5.73, 1.90)
            tf = card.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_right = Inches(0.25)
            tf.margin_top = Inches(0.18)

            p_num = tf.paragraphs[0]
            r_num = p_num.add_run()
            r_num.text = num
            r_num.font.name = FONT_DISPLAY
            r_num.font.size = Pt(38)
            r_num.font.bold = True
            r_num.font.color.rgb = COLOR_LIME

            p_title = tf.add_paragraph()
            p_title.space_before = Pt(2)
            r_title = p_title.add_run()
            r_title.text = title
            r_title.font.name = FONT_BODY
            r_title.font.size = Pt(10.5)
            r_title.font.bold = True
            r_title.font.color.rgb = COLOR_CYAN

            p_desc = tf.add_paragraph()
            p_desc.space_before = Pt(4)
            r_desc = p_desc.add_run()
            r_desc.text = desc
            r_desc.font.name = FONT_BODY
            r_desc.font.size = Pt(9.5)
            r_desc.font.color.rgb = COLOR_TEXT_MUTED

        self.add_callout(
            slide,
            0.75, 6.38, 11.83, 0.48,
            "📈 Eficiência Institucional: Eliminação total de reuniões de alinhamento de números; o painel é a fonte soberana da verdade.",
            accent_color=COLOR_LIME,
        )
        self.add_footer(slide)

    # ==============================================================================
    # SLIDE 5: RAIO-X DAS MÉTRICAS & METODOLOGIA DE EVIDÊNCIAS
    # ==============================================================================
    def build_slide_5(self):
        slide = self.add_slide()
        self.add_header(
            slide,
            tag_text="AUDITORIA DE INDICADORES  /  METODOLOGIA & EVIDÊNCIAS",
            title_text="Raio-X das métricas: base de cálculo e evidências operacionais.",
            subtitle_text="Como cada recurso e métrica do JiraView é validado em tempo real contra as bases corporativas.",
            slide_num_str="05",
            tag_color=COLOR_CYAN,
        )

        evidencias = [
            (0.75, 2.25, "⚡ DECOMPOSIÇÃO DO BACKLOG ATIVO (25 ATIVAS / 25 CONCLUÍDAS)",
             "ESCOPO: Carteira ativa no ciclo contratual de 90 dias.  |  EVIDÊNCIA: 0 em Triagem + 24 em Atendimento Técnico egSYS + 1 Aguardando Validação do Cliente. 25 itens entregues segregados no histórico (total 50)."),
            (0.75, 3.30, "🎯 TRILHA DA JORNADA & TEMPO ACUMULADO POR FASE",
             "ESCOPO: Radiografia profunda do ciclo de vida no endpoint /journey.  |  EVIDÊNCIA: Parser do histórico de transições do Jira Cloud calculando dias e horas exatos de permanência em cada uma das 8 fases."),
            (0.75, 4.35, "📊 DASHBOARDS GRÁFICOS & VENDORING LOCAL (204KB)",
             "ESCOPO: Renderização gráfica estável em qualquer rede corporativa restrita.  |  EVIDÊNCIA: Chart.js empacotado localmente no backend; gráficos de barras, donut, polar e linha carregam 100% offline."),
            (0.75, 5.40, "💾 PERSISTÊNCIA DE FILTROS SALVOS & EXPORTAÇÃO (# 1 A N)",
             "ESCOPO: Autonomia de personalização de relatórios pelo próprio gestor.  |  EVIDÊNCIA: Tabela de filtros com permissão 0600 em data/filtros.json e módulo de download em Excel nativo (.xlsx), MD e CSV."),
        ]

        for left, top, tag_title, content in evidencias:
            card = self.add_card(slide, left, top, 11.83, 0.95)
            tf = card.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_right = Inches(0.25)
            tf.margin_top = Inches(0.12)

            p1 = tf.paragraphs[0]
            r1 = p1.add_run()
            r1.text = tag_title
            r1.font.name = FONT_DISPLAY
            r1.font.size = Pt(11)
            r1.font.bold = True
            r1.font.color.rgb = COLOR_LIME

            p2 = tf.add_paragraph()
            p2.space_before = Pt(3)
            r2 = p2.add_run()
            r2.text = content
            r2.font.name = FONT_BODY
            r2.font.size = Pt(9.5)
            r2.font.color.rgb = COLOR_TEXT_LIGHT

        self.add_callout(
            slide,
            0.75, 6.45, 11.83, 0.42,
            "🔍 Critério de Rigor Técnico: Dados auditados diretamente no backend FastAPI, deploy/estados.yaml e banco SQLite auth.db.",
            accent_color=COLOR_CYAN,
        )
        self.add_footer(slide)

    # ==============================================================================
    # SLIDE 6: A ESTEIRA DE 8 ETAPAS & AS NOVAS FUNCIONALIDADES DO PAINEL
    # ==============================================================================
    def build_slide_6(self):
        slide = self.add_slide()
        self.add_header(
            slide,
            tag_text="CATÁLOGO DE MELHORIAS  /  FUNCIONALIDADES DO PAINEL",
            title_text="A jornada na esteira de ação e os recursos de observabilidade.",
            subtitle_text="Todas as ferramentas que o cliente precisava e que o portal padrão do Jira nunca ofereceu.",
            slide_num_str="06",
            tag_color=COLOR_CYAN,
        )

        # Painel Esquerdo: A Esteira de Ação e Posse da Bola
        card_l = self.add_card(slide, 0.75, 2.25, 5.73, 4.55)
        tf_l = card_l.text_frame
        tf_l.word_wrap = True
        tf_l.margin_left = tf_l.margin_right = Inches(0.25)
        tf_l.margin_top = Inches(0.20)

        p_lt = tf_l.paragraphs[0]
        r_lt = p_lt.add_run()
        r_lt.text = "A TRILHA DO CAMINHO (ESTEIRA DE 8 ETAPAS)"
        r_lt.font.name = FONT_DISPLAY
        r_lt.font.size = Pt(13)
        r_lt.font.bold = True
        r_lt.font.color.rgb = COLOR_CYAN

        p_lsub = tf_l.add_paragraph()
        r_lsub = p_lsub.add_run()
        r_lsub.text = "Ciclo de vida transparente com indicação de Posse da Bola:"
        r_lsub.font.name = FONT_BODY
        r_lsub.font.size = Pt(9.5)
        r_lsub.font.color.rgb = COLOR_TEXT_MUTED

        etapas = [
            ("01. Triagem (N1)", "Recepção, qualificação inicial e abertura de chamado."),
            ("02. Triagem (N2)", "Diagnóstico técnico avançado e reprodução de logs."),
            ("03. Análise de Dev", "Refinamento arquitetural e desenho de solução técnica."),
            ("04. Em Desenvolvimento", "Construção de código ativo e testes de unidade."),
            ("05. QA & Testes", "Homologação técnica e baterias de testes automatizados."),
            ("06. Validação Interna N1", "Conferência final pelo analista antes da entrega."),
            ("07. Homologação Cliente", "AÇÃO DO CLIENTE: Testes em staging e aceite formal."),
            ("08. Concluído", "FINALIZADO: Liberação em produção e encerramento."),
        ]

        for num_nome, desc in etapas:
            p = tf_l.add_paragraph()
            p.space_before = Pt(3)
            r1 = p.add_run()
            r1.text = "• " + num_nome + " ➔ "
            r1.font.name = FONT_BODY
            r1.font.size = Pt(9.5)
            r1.font.bold = True
            r1.font.color.rgb = COLOR_LIME if "07" in num_nome or "08" in num_nome else COLOR_TEXT_LIGHT

            r2 = p.add_run()
            r2.text = desc
            r2.font.name = FONT_BODY
            r2.font.size = Pt(9.0)
            r2.font.color.rgb = COLOR_TEXT_MUTED

        # Painel Direito: Lista Completa de Melhorias Funcionais
        card_r = self.add_card(slide, 6.85, 2.25, 5.73, 4.55)
        tf_r = card_r.text_frame
        tf_r.word_wrap = True
        tf_r.margin_left = tf_r.margin_right = Inches(0.25)
        tf_r.margin_top = Inches(0.20)

        p_rt = tf_r.paragraphs[0]
        r_rt = p_rt.add_run()
        r_rt.text = "FUNCIONALIDADES EXCLUSIVAS DO JIRAVIEW"
        r_rt.font.name = FONT_DISPLAY
        r_rt.font.size = Pt(13)
        r_rt.font.bold = True
        r_rt.font.color.rgb = COLOR_CYAN

        p_rsub = tf_r.add_paragraph()
        r_rsub = p_rsub.add_run()
        r_rsub.text = "Recursos de observabilidade ausentes no portal legado:"
        r_rsub.font.name = FONT_BODY
        r_rsub.font.size = Pt(9.5)
        r_rsub.font.color.rgb = COLOR_TEXT_MUTED

        melhorias = [
            ("Gaveta Lateral de Diagnóstico (Drawer): ", "Abre em 1 clique sem reload, exibindo a esteira, o tempo exato acumulado em cada etapa e a posse da bola."),
            ("Módulo de Filtros Salvos pelo Gestor: ", "Permite criar filtros personalizados e salvá-los no perfil (/filtros) para reuso diário com um clique."),
            ("Dashboard com 4 Gráficos Interativos: ", "Gráficos Chart.js (status, prioridade, solicitantes e linha temporal) atualizados em tempo real."),
            ("Rastreabilidade de Engenharia (issuelinks): ", "Vínculo visual direto com tarefas de desenvolvimento (PSC) ligadas ao chamado de atendimento."),
            ("Exportação Corporativa (# 1 a N): ", "Download imediato em Excel (.xlsx), Markdown e CSV com índice sequencial (#) para pautas de reunião."),
        ]

        for m_title, m_desc in melhorias:
            p = tf_r.add_paragraph()
            p.space_before = Pt(6)
            r1 = p.add_run()
            r1.text = "• " + m_title
            r1.font.name = FONT_BODY
            r1.font.size = Pt(9.5)
            r1.font.bold = True
            r1.font.color.rgb = COLOR_TEXT_LIGHT

            r2 = p.add_run()
            r2.text = m_desc
            r2.font.name = FONT_BODY
            r2.font.size = Pt(9.0)
            r2.font.color.rgb = COLOR_TEXT_MUTED

        self.add_footer(slide)

    # ==============================================================================
    # SLIDE 7: ESCALA MULTI-ESTADO & ARQUITETURA POR CONFIGURAÇÃO
    # ==============================================================================
    def build_slide_7(self):
        slide = self.add_slide()
        self.add_header(
            slide,
            tag_text="ESCALA MULTI-ESTADO  /  CONFIGURAÇÃO DECLARATIVA",
            title_text="O modelo multi-estado: expansão nacional sem custo de infra.",
            subtitle_text="Um único produto atende múltiplos clientes públicos mantendo isolamento absoluto de dados e permissões.",
            slide_num_str="07",
            tag_color=COLOR_CYAN,
        )

        # Painel Esquerdo: Core e Estados
        card_l = self.add_card(slide, 0.75, 2.25, 7.50, 4.55)
        tf_l = card_l.text_frame
        tf_l.word_wrap = True
        tf_l.margin_left = tf_l.margin_right = Inches(0.25)
        tf_l.margin_top = Inches(0.20)

        p_lt = tf_l.paragraphs[0]
        r_lt = p_lt.add_run()
        r_lt.text = "JIRAVIEW CORE — NÚCLEO INTELIGENTE MULTI-ESTADO"
        r_lt.font.name = FONT_DISPLAY
        r_lt.font.size = Pt(13)
        r_lt.font.bold = True
        r_lt.font.color.rgb = COLOR_CYAN

        p_lsub = tf_l.add_paragraph()
        r_lsub = p_lsub.add_run()
        r_lsub.text = "Topologia declarativa onde o estado é escopo lógico de dados:"
        r_lsub.font.name = FONT_BODY
        r_lsub.font.size = Pt(9.5)
        r_lsub.font.color.rgb = COLOR_TEXT_MUTED

        estados_info = [
            ("🟢 Santa Catarina (PMSC)", "Projeto HDPMSC • Gestores João Mário, Ilclemar, Alex e Cap Thiesen • 4 áreas (Cidadão, SADE, Integração, Operações)."),
            ("🟢 Tocantins (PMTO)", "Projeto HDPMTO • Configurado e pronto para ativação declarativa sem novas instâncias."),
            ("🟢 Amazonas (PMAM) & Rondônia (PMRO)", "Projetos HDSUPAM e HDRO • Suporte operacional mapeado para atendimento integrado."),
            ("🟢 Paraná (PMPR) & Guardas Municipais", "Projetos HDPMPR e HDGM • Expansão para suporte municipal e metropolitano."),
            ("🟢 Coordenação de Suporte Geral", "Visão panorâmica dos 73 espaços Jira com busca universal e gestão de usuários."),
        ]

        for est_title, est_desc in estados_info:
            p = tf_l.add_paragraph()
            p.space_before = Pt(7)
            r1 = p.add_run()
            r1.text = est_title + "\n"
            r1.font.name = FONT_BODY
            r1.font.size = Pt(10)
            r1.font.bold = True
            r1.font.color.rgb = COLOR_LIME

            r2 = p.add_run()
            r2.text = est_desc
            r2.font.name = FONT_BODY
            r2.font.size = Pt(9.0)
            r2.font.color.rgb = COLOR_TEXT_MUTED

        # Painel Direito Superior: Onboarding Rápido
        card_ru = self.add_card(slide, 8.55, 2.25, 4.03, 2.15)
        tf_ru = card_ru.text_frame
        tf_ru.word_wrap = True
        tf_ru.margin_left = tf_ru.margin_right = Inches(0.20)
        tf_ru.margin_top = Inches(0.18)

        p_ru1 = tf_ru.paragraphs[0]
        r_ru1 = p_ru1.add_run()
        r_ru1.text = "< 15 min"
        r_ru1.font.name = FONT_DISPLAY
        r_ru1.font.size = Pt(32)
        r_ru1.font.bold = True
        r_ru1.font.color.rgb = COLOR_LIME

        p_ru2 = tf_ru.add_paragraph()
        r_ru2 = p_ru2.add_run()
        r_ru2.text = "ONBOARDING DE NOVO ESTADO"
        r_ru2.font.name = FONT_BODY
        r_ru2.font.size = Pt(10)
        r_ru2.font.bold = True
        r_ru2.font.color.rgb = COLOR_CYAN

        p_ru3 = tf_ru.add_paragraph()
        p_ru3.space_before = Pt(3)
        r_ru3 = p_ru3.add_run()
        r_ru3.text = "Basta incluir um bloco no deploy/estados.yaml e cadastrar os gestores no painel administrativo. Sem novos containers."
        r_ru3.font.name = FONT_BODY
        r_ru3.font.size = Pt(9.0)
        r_ru3.font.color.rgb = COLOR_TEXT_MUTED

        # Painel Direito Inferior: Contenção Cgroups
        card_rd = self.add_card(slide, 8.55, 4.65, 4.03, 2.15)
        tf_rd = card_rd.text_frame
        tf_rd.word_wrap = True
        tf_rd.margin_left = tf_rd.margin_right = Inches(0.20)
        tf_rd.margin_top = Inches(0.18)

        p_rd1 = tf_rd.paragraphs[0]
        r_rd1 = p_rd1.add_run()
        r_rd1.text = "512 MB"
        r_rd1.font.name = FONT_DISPLAY
        r_rd1.font.size = Pt(32)
        r_rd1.font.bold = True
        r_rd1.font.color.rgb = COLOR_LIME

        p_rd2 = tf_rd.add_paragraph()
        r_rd2 = p_rd2.add_run()
        r_rd2.text = "CONTENÇÃO MÁXIMA CGROUPS"
        r_rd2.font.name = FONT_BODY
        r_rd2.font.size = Pt(10)
        r_rd2.font.bold = True
        r_rd2.font.color.rgb = COLOR_CYAN

        p_rd3 = tf_rd.add_paragraph()
        p_rd3.space_before = Pt(3)
        r_rd3 = p_rd3.add_run()
        r_rd3.text = "Consumo ultra-leve com 1.0 CPU e 150 PIDs de limite. Zero impacto no servidor de monitoramento e zero vazamento de recursos."
        r_rd3.font.name = FONT_BODY
        r_rd3.font.size = Pt(9.0)
        r_rd3.font.color.rgb = COLOR_TEXT_MUTED

        self.add_footer(slide)

    # ==============================================================================
    # SLIDE 8: ROADMAP ESTRATÉGICO & CONCLUSÃO EXECUTIVA
    # ==============================================================================
    def build_slide_8(self):
        slide = self.add_slide()
        self.add_header(
            slide,
            tag_text="EVOLUÇÃO CONTÍNUA  /  EXECUTIVE TAKEAWAY",
            title_text="O JiraView consolida o padrão executivo de transparência.",
            subtitle_text="Da eliminação de atritos de atendimento à gestão inteligente orientada a evidências e dados reais.",
            slide_num_str="08",
            tag_color=COLOR_CYAN,
        )

        colunas = [
            (0.75, 2.25, 3.75, 4.55, "ENTREGUE EM 2026 (v0.4.1)", "Operação Consolidada", [
                ("Banner de Síntese Matemática: ", "Decomposição exata com 25 tarefas abertas e 25 concluídas (ciclo 90d)."),
                ("Esteira Canônica de 8 Etapas: ", "Trilha visual da jornada e indicação de Posse da Bola."),
                ("Dashboards & Gráficos Táticos: ", "4 visões visuais integradas para o cliente."),
                ("Filtros Salvos pelo Gestor: ", "Módulo CRUD (/filtros) para salvar visões favoritas."),
                ("Exportação Corporativa (# 1 a N): ", "Planilhas Excel (.xlsx), Markdown e CSV com índice."),
            ]),
            (4.79, 2.25, 3.75, 4.55, "ROADMAP 2026 / 2027", "Próximas Inovações (v0.5 / v1.0)", [
                ("Ativação TO, AM e RO: ", "Expansão dos novos estados via configuração declarativa."),
                ("Alertas Proativos de SLA: ", "Webhooks em tempo real notificando chamados em risco."),
                ("Auditoria Forense de Acessos: ", "Trilha completa de acessos e ações dos gestores em banco."),
                ("Assistente com IA egSYS: ", "Sumarização executiva automática de chamados complexos."),
                ("App Mobile Nativo PWA: ", "Acesso ultrarrápido para gestores em tablets e smartphones."),
            ]),
            (8.83, 2.25, 3.75, 4.55, "SÍNTESE DE VALOR", "Retorno Estratégico para Diretoria", [
                ("⚡ VELOCIDADE: ", "Diagnóstico executivo de chamados em ≤ 3 segundos sem sobrecarga."),
                ("🛡️ CONTROLE: ", "Isolamento RBAC por estado, Cgroups 512MB e filtros salvos."),
                ("🤝 CONFIANÇA: ", "Fim das discussões de backlog ativo através de dados matemáticos."),
                ("💰 EFICIÊNCIA: ", "Zero infraestrutura adicional para expansão em novos contratos."),
                ("🏛️ SOBERANIA: ", "Tecnologia própria egSYS, auditável e independente de licenças terceiras."),
            ]),
        ]

        for left, top, width, height, tag, title, bullets in colunas:
            card = self.add_card(slide, left, top, width, height)
            tf = card.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_right = Inches(0.25)
            tf.margin_top = Inches(0.22)

            p1 = tf.paragraphs[0]
            r1 = p1.add_run()
            r1.text = tag
            r1.font.name = FONT_BODY
            r1.font.size = Pt(10)
            r1.font.bold = True
            r1.font.color.rgb = COLOR_CYAN

            p2 = tf.add_paragraph()
            p2.space_before = Pt(3)
            r2 = p2.add_run()
            r2.text = title
            r2.font.name = FONT_DISPLAY
            r2.font.size = Pt(13)
            r2.font.bold = True
            r2.font.color.rgb = COLOR_TEXT_LIGHT

            for b_label, b_text in bullets:
                p = tf.add_paragraph()
                p.space_before = Pt(6)
                rb1 = p.add_run()
                rb1.text = "• " + b_label
                rb1.font.name = FONT_BODY
                rb1.font.size = Pt(9.5)
                rb1.font.bold = True
                rb1.font.color.rgb = COLOR_LIME if "VELOCIDADE" in b_label or "CONFIANÇA" in b_label or "EFICIÊNCIA" in b_label else COLOR_TEXT_LIGHT

                rb2 = p.add_run()
                rb2.text = b_text
                rb2.font.name = FONT_BODY
                rb2.font.size = Pt(9.0)
                rb2.font.color.rgb = COLOR_TEXT_MUTED

        self.add_footer(slide)

    def generate_all(self):
        print("Montando Slide 1 (Capa & Executive Brief)...")
        self.build_slide_1()
        print("Montando Slide 2 (O Desafio & Riscos)...")
        self.build_slide_2()
        print("Montando Slide 3 (A Virada Arquitetural)...")
        self.build_slide_3()
        print("Montando Slide 4 (Impacto Mensurado & ROI)...")
        self.build_slide_4()
        print("Montando Slide 5 (Raio-X das Métricas)...")
        self.build_slide_5()
        print("Montando Slide 6 (A Esteira de 8 Etapas & Funcionalidades)...")
        self.build_slide_6()
        print("Montando Slide 7 (Escala Multi-Estado)...")
        self.build_slide_7()
        print("Montando Slide 8 (Roadmap & Conclusão)...")
        self.build_slide_8()

        self.prs.save(self.output_pptx)
        print(f"✅ PPTX salvo com sucesso em: {self.output_pptx}")

        # Exportar para PDF via LibreOffice headless
        out_dir = os.path.dirname(os.path.abspath(self.output_pptx)) or "."
        try:
            print(f"Executando conversão headless para PDF via LibreOffice...")
            subprocess.run(
                ["libreoffice", "--headless", "--convert-to", "pdf", self.output_pptx, "--outdir", out_dir],
                check=True,
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE
            )
            pdf_path = os.path.splitext(self.output_pptx)[0] + ".pdf"
            print(f"✅ PDF gerado com sucesso em: {pdf_path}")
        except Exception as e:
            print(f"⚠️ Aviso na conversão LibreOffice para PDF: {e}")


if __name__ == "__main__":
    target_pptx = sys.argv[1] if len(sys.argv) > 1 else "/home/prado/Documents/egSYS-JiraView-CEO-Executivo.pptx"
    builder = JiraViewDeckBuilder(target_pptx)
    builder.generate_all()
